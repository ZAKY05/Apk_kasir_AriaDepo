"""
Customers View - Halaman manajemen pelanggan
Menampilkan data pelanggan beserta nominal hutang
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from styles import COLORS, FONTS, ICONS, format_currency
import database as db


class CustomersView(tk.Frame):
    def __init__(self, parent, user):
        super().__init__(parent, bg=COLORS['bg_light'])
        self.user = user
        self._build_ui()
        self.refresh_data()

    def _build_ui(self):
        # Header
        header = tk.Frame(self, bg=COLORS['bg_light'])
        header.pack(fill='x', padx=25, pady=(25, 15))

        tk.Label(
            header, text=f"{ICONS['customer']}  Data Pelanggan",
            font=FONTS['title'], bg=COLORS['bg_light'], fg=COLORS['text_dark']
        ).pack(side='left')

        add_btn = tk.Button(
            header, text=f"  {ICONS['add']}  Tambah Pelanggan  ", font=FONTS['button'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._show_add_dialog, pady=8
        )
        add_btn.pack(side='right')
        add_btn.bind('<Enter>', lambda e: add_btn.config(bg=COLORS['primary_hover']))
        add_btn.bind('<Leave>', lambda e: add_btn.config(bg=COLORS['primary']))

        # Summary cards
        summary_frame = tk.Frame(self, bg=COLORS['bg_light'])
        summary_frame.pack(fill='x', padx=25, pady=(0, 15))
        summary_frame.columnconfigure((0, 1, 2), weight=1)

        self.total_card = tk.Frame(summary_frame, bg=COLORS['white'], padx=20, pady=15,
                                    highlightbackground=COLORS['border_light'], highlightthickness=1)
        self.total_card.grid(row=0, column=0, sticky='nsew', padx=(0, 8))
        tk.Label(self.total_card, text="Total Pelanggan", font=FONTS['small_bold'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(anchor='w')
        self.total_customer_label = tk.Label(self.total_card, text="0", font=FONTS['medium_number'],
                                              bg=COLORS['white'], fg=COLORS['primary'])
        self.total_customer_label.pack(anchor='w')

        self.debt_card = tk.Frame(summary_frame, bg=COLORS['white'], padx=20, pady=15,
                                   highlightbackground=COLORS['border_light'], highlightthickness=1)
        self.debt_card.grid(row=0, column=1, sticky='nsew', padx=4)
        tk.Label(self.debt_card, text="Total Hutang Pelanggan", font=FONTS['small_bold'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(anchor='w')
        self.total_debt_label = tk.Label(self.debt_card, text="Rp 0", font=FONTS['medium_number'],
                                          bg=COLORS['white'], fg=COLORS['danger'])
        self.total_debt_label.pack(anchor='w')

        self.debtor_card = tk.Frame(summary_frame, bg=COLORS['white'], padx=20, pady=15,
                                     highlightbackground=COLORS['border_light'], highlightthickness=1)
        self.debtor_card.grid(row=0, column=2, sticky='nsew', padx=(8, 0))
        tk.Label(self.debtor_card, text="Pelanggan Berhutang", font=FONTS['small_bold'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(anchor='w')
        self.debtor_count_label = tk.Label(self.debtor_card, text="0", font=FONTS['medium_number'],
                                            bg=COLORS['white'], fg=COLORS['warning'])
        self.debtor_count_label.pack(anchor='w')

        # Filter & search
        filter_frame = tk.Frame(self, bg=COLORS['white'], padx=15, pady=10,
                                 highlightbackground=COLORS['border_light'], highlightthickness=1)
        filter_frame.pack(fill='x', padx=25, pady=(0, 15))

        tk.Label(filter_frame, text=ICONS['search'], font=('Segoe UI', 14),
                 bg=COLORS['white']).pack(side='left', padx=(0, 10))

        self.search_var = tk.StringVar()
        self.search_var.trace('w', lambda *a: self.refresh_data())
        tk.Entry(filter_frame, textvariable=self.search_var, font=FONTS['input'],
                 relief='flat', bg=COLORS['white'], width=30,
                 insertbackground=COLORS['primary']).pack(side='left', fill='x', expand=True)

        # Filter tombol
        self.filter_var = tk.StringVar(value="all")

        filter_all = tk.Radiobutton(
            filter_frame, text="Semua", variable=self.filter_var, value="all",
            font=FONTS['body'], bg=COLORS['white'], fg=COLORS['text_dark'],
            selectcolor=COLORS['white'], activebackground=COLORS['white'],
            command=self.refresh_data, cursor='hand2'
        )
        filter_all.pack(side='right', padx=5)

        filter_debt = tk.Radiobutton(
            filter_frame, text="🔴 Berhutang", variable=self.filter_var, value="debt",
            font=FONTS['body_bold'], bg=COLORS['white'], fg=COLORS['danger'],
            selectcolor=COLORS['white'], activebackground=COLORS['white'],
            command=self.refresh_data, cursor='hand2'
        )
        filter_debt.pack(side='right', padx=5)

        filter_clear = tk.Radiobutton(
            filter_frame, text="🟢 Lunas", variable=self.filter_var, value="clear",
            font=FONTS['body'], bg=COLORS['white'], fg=COLORS['success'],
            selectcolor=COLORS['white'], activebackground=COLORS['white'],
            command=self.refresh_data, cursor='hand2'
        )
        filter_clear.pack(side='right', padx=5)

        # ===== Action frame - pack di BOTTOM dulu =====
        action_frame = tk.Frame(self, bg=COLORS['bg_light'], pady=5)
        action_frame.pack(side='bottom', fill='x', padx=25, pady=(5, 15))

        edit_btn = tk.Button(
            action_frame, text=f"  {ICONS['edit']}  Edit  ", font=FONTS['body_bold'],
            bg=COLORS['primary_light'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._show_edit_dialog, pady=8, padx=12
        )
        edit_btn.pack(side='left', padx=(0, 8))
        edit_btn.bind('<Enter>', lambda e: edit_btn.config(bg=COLORS['primary']))
        edit_btn.bind('<Leave>', lambda e: edit_btn.config(bg=COLORS['primary_light']))

        pay_btn = tk.Button(
            action_frame, text=f"  {ICONS['payment']}  Bayar Hutang  ", font=FONTS['body_bold'],
            bg=COLORS['success'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._show_pay_debt_dialog, pady=8, padx=12
        )
        pay_btn.pack(side='left', padx=(0, 8))
        pay_btn.bind('<Enter>', lambda e: pay_btn.config(bg='#1B5E20'))
        pay_btn.bind('<Leave>', lambda e: pay_btn.config(bg=COLORS['success']))

        history_btn = tk.Button(
            action_frame, text=f"  {ICONS['debt']}  Riwayat Hutang  ", font=FONTS['body_bold'],
            bg=COLORS['warning'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._show_debt_history, pady=8, padx=12
        )
        history_btn.pack(side='left', padx=(0, 8))
        history_btn.bind('<Enter>', lambda e: history_btn.config(bg='#E65100'))
        history_btn.bind('<Leave>', lambda e: history_btn.config(bg=COLORS['warning']))

        del_btn = tk.Button(
            action_frame, text=f"  {ICONS['delete']}  Hapus  ", font=FONTS['body_bold'],
            bg=COLORS['danger'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._delete_customer, pady=8, padx=12
        )
        del_btn.pack(side='left')
        del_btn.bind('<Enter>', lambda e: del_btn.config(bg='#B71C1C'))
        del_btn.bind('<Leave>', lambda e: del_btn.config(bg=COLORS['danger']))

        self.count_label = tk.Label(
            action_frame, text="", font=FONTS['small'],
            bg=COLORS['bg_light'], fg=COLORS['text_light']
        )
        self.count_label.pack(side='right')

        # ===== Tabel pelanggan - pack SETELAH action_frame =====
        table_frame = tk.Frame(self, bg=COLORS['white'],
                                highlightbackground=COLORS['border_light'], highlightthickness=1)
        table_frame.pack(fill='both', expand=True, padx=25, pady=(0, 5))

        style = ttk.Style()
        style.configure('Customers.Treeview',
                         background=COLORS['white'],
                         foreground=COLORS['text_dark'],
                         fieldbackground=COLORS['white'],
                         font=FONTS['table'],
                         rowheight=38)
        style.configure('Customers.Treeview.Heading',
                         background=COLORS['primary'],
                         foreground=COLORS['text_white'],
                         font=FONTS['table_header'],
                         relief='flat')
        style.map('Customers.Treeview.Heading',
                   background=[('active', COLORS['primary_hover'])])
        style.map('Customers.Treeview',
                   background=[('selected', COLORS['accent_light'])],
                   foreground=[('selected', COLORS['text_dark'])])

        columns = ('id', 'name', 'phone', 'address', 'debt', 'status')
        self.tree = ttk.Treeview(table_frame, columns=columns, show='headings',
                                  style='Customers.Treeview', selectmode='browse')

        self.tree.heading('id', text='ID')
        self.tree.heading('name', text='Nama Pelanggan')
        self.tree.heading('phone', text='No. Telepon')
        self.tree.heading('address', text='Alamat')
        self.tree.heading('debt', text='Hutang')
        self.tree.heading('status', text='Status')

        self.tree.column('id', width=50, anchor='center')
        self.tree.column('name', width=180, anchor='w')
        self.tree.column('phone', width=130, anchor='center')
        self.tree.column('address', width=200, anchor='w')
        self.tree.column('debt', width=140, anchor='e')
        self.tree.column('status', width=100, anchor='center')

        scrollbar = ttk.Scrollbar(table_frame, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=scrollbar.set)
        self.tree.pack(side='left', fill='both', expand=True)
        scrollbar.pack(side='right', fill='y')

        self.tree.tag_configure('has_debt', background=COLORS['danger_bg'])
        self.tree.tag_configure('no_debt', background=COLORS['white'])

        self.tree.bind('<Double-1>', lambda e: self._show_edit_dialog())

    def refresh_data(self):
        """Refresh data tabel pelanggan"""
        for item in self.tree.get_children():
            self.tree.delete(item)

        customers = db.get_all_customers()
        search = self.search_var.get().lower()
        filter_type = self.filter_var.get()

        count = 0

        for c in customers:
            if search and search not in c['name'].lower() and search not in (c['phone'] or '').lower():
                continue

            has_debt = c['debt_amount'] > 0
            if filter_type == "debt" and not has_debt:
                continue
            if filter_type == "clear" and has_debt:
                continue

            status = "BERHUTANG" if has_debt else "LUNAS"
            tags = ('has_debt',) if has_debt else ('no_debt',)

            self.tree.insert('', 'end', values=(
                c['id'], c['name'], c['phone'] or '-',
                c['address'] or '-',
                format_currency(c['debt_amount']),
                status
            ), tags=tags)
            count += 1

        # Update summary cards
        all_customers = db.get_all_customers()
        all_total_debt = sum(c['debt_amount'] for c in all_customers)
        all_debtor_count = sum(1 for c in all_customers if c['debt_amount'] > 0)

        self.total_customer_label.config(text=str(len(all_customers)))
        self.total_debt_label.config(text=format_currency(all_total_debt))
        self.debtor_count_label.config(text=str(all_debtor_count))
        self.count_label.config(text=f"Menampilkan: {count} pelanggan")

    def _show_add_dialog(self):
        self._show_customer_dialog("Tambah Pelanggan Baru")

    def _show_edit_dialog(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pelanggan yang ingin diedit!")
            return
        values = self.tree.item(selected[0])['values']
        self._show_customer_dialog("Edit Pelanggan", values)

    def _show_customer_dialog(self, title, values=None):
        dialog = tk.Toplevel(self)
        dialog.title(title)
        dialog.geometry("460x440")
        dialog.minsize(420, 360)
        dialog.resizable(True, True)
        dialog.configure(bg=COLORS['white'])
        dialog.transient(self)
        dialog.grab_set()

        dialog.update_idletasks()
        x = (dialog.winfo_screenwidth() // 2) - 230
        y = (dialog.winfo_screenheight() // 2) - 220
        dialog.geometry(f'+{x}+{y}')

        # Header
        tk.Label(dialog, text=title, font=FONTS['subtitle'],
                 bg=COLORS['white'], fg=COLORS['primary_dark']).pack(side='top', pady=(15, 10))

        sep = tk.Frame(dialog, bg=COLORS['accent_light'], height=2)
        sep.pack(side='top', fill='x', padx=30)

        # Buttons packed at BOTTOM first so they are ALWAYS visible
        btn_frame = tk.Frame(dialog, bg=COLORS['white'])
        btn_frame.pack(side='bottom', fill='x', padx=30, pady=15)

        form = tk.Frame(dialog, bg=COLORS['white'], padx=30)
        form.pack(side='top', fill='both', expand=True, pady=10)

        fields = {}

        def add_field(label_text, key, default=""):
            tk.Label(form, text=label_text, font=FONTS['body_bold'],
                     bg=COLORS['white'], fg=COLORS['text_dark'], anchor='w').pack(fill='x', pady=(6, 2))
            var = tk.StringVar(value=str(default))
            entry = tk.Entry(form, textvariable=var, font=FONTS['input'],
                              relief='solid', bd=1, bg=COLORS['bg_light'])
            entry.pack(fill='x', ipady=5)
            entry.configure(highlightthickness=1, highlightcolor=COLORS['primary'],
                             highlightbackground=COLORS['border'])
            fields[key] = var

        add_field("Nama Pelanggan", 'name', values[1] if values else "")
        add_field("No. Telepon", 'phone', values[2] if values and values[2] != '-' else "")
        add_field("Alamat", 'address', values[3] if values and values[3] != '-' else "")

        def save():
            name = fields['name'].get().strip()
            phone = fields['phone'].get().strip()
            address = fields['address'].get().strip()

            if not name:
                messagebox.showwarning("Peringatan", "Nama pelanggan harus diisi!", parent=dialog)
                return

            if values:
                db.update_customer(values[0], name, phone, address)
                messagebox.showinfo("Berhasil", "Data pelanggan berhasil diupdate!", parent=dialog)
            else:
                db.add_customer(name, phone, address)
                messagebox.showinfo("Berhasil", "Pelanggan berhasil ditambahkan!", parent=dialog)

            dialog.destroy()
            self.refresh_data()

        save_btn = tk.Button(
            btn_frame, text="  💾  Simpan  ", font=FONTS['button'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=save, padx=25, pady=8
        )
        save_btn.pack(side='left', padx=(0, 10))
        save_btn.bind('<Enter>', lambda e: save_btn.config(bg=COLORS['primary_hover']))
        save_btn.bind('<Leave>', lambda e: save_btn.config(bg=COLORS['primary']))

        cancel_btn = tk.Button(
            btn_frame, text="  Batal  ", font=FONTS['button'],
            bg=COLORS['border'], fg=COLORS['text_dark'], relief='flat',
            cursor='hand2', command=dialog.destroy, padx=20, pady=8
        )
        cancel_btn.pack(side='left')

    def _show_pay_debt_dialog(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pelanggan!")
            return

        values = self.tree.item(selected[0])['values']
        customer_id = values[0]
        customer_name = values[1]
        customer = db.get_customer_by_id(customer_id)

        if customer['debt_amount'] <= 0:
            messagebox.showinfo("Info", f"{customer_name} tidak memiliki hutang!")
            return

        dialog = tk.Toplevel(self)
        dialog.title(f"Bayar Hutang - {customer_name}")
        dialog.geometry("450x480")
        dialog.minsize(400, 400)
        dialog.resizable(True, True)
        dialog.configure(bg=COLORS['white'])
        dialog.transient(self)
        dialog.grab_set()

        dialog.update_idletasks()
        x = (dialog.winfo_screenwidth() // 2) - 225
        y = (dialog.winfo_screenheight() // 2) - 240
        dialog.geometry(f'+{x}+{y}')

        # Header
        tk.Label(dialog, text="Bayar Hutang", font=FONTS['subtitle'],
                 bg=COLORS['white'], fg=COLORS['primary_dark']).pack(side='top', pady=(15, 8))

        info_frame = tk.Frame(dialog, bg=COLORS['info_bg'], padx=15, pady=10)
        info_frame.pack(side='top', fill='x', padx=30, pady=(0, 10))

        tk.Label(info_frame, text=f"Pelanggan: {customer_name}", font=FONTS['body_bold'],
                 bg=COLORS['info_bg'], fg=COLORS['text_dark']).pack(anchor='w')
        tk.Label(info_frame, text=f"Total Hutang: {format_currency(customer['debt_amount'])}",
                 font=FONTS['body_bold'], bg=COLORS['info_bg'], fg=COLORS['danger']).pack(anchor='w')

        # Buttons packed at BOTTOM first so they are ALWAYS visible
        btn_frame = tk.Frame(dialog, bg=COLORS['white'])
        btn_frame.pack(side='bottom', fill='x', padx=30, pady=15)

        form = tk.Frame(dialog, bg=COLORS['white'], padx=30)
        form.pack(side='top', fill='both', expand=True)

        tk.Label(form, text="Jumlah Bayar (Rp)", font=FONTS['body_bold'],
                 bg=COLORS['white'], fg=COLORS['text_dark']).pack(anchor='w', pady=(3, 2))
        amount_var = tk.StringVar()
        amount_entry = tk.Entry(form, textvariable=amount_var, font=FONTS['input'],
                                 relief='solid', bd=1, bg=COLORS['bg_light'])
        amount_entry.pack(fill='x', ipady=6)
        amount_entry.configure(highlightthickness=1, highlightcolor=COLORS['primary'],
                                highlightbackground=COLORS['border'])
        amount_entry.focus_set()

        tk.Label(form, text="Catatan", font=FONTS['body_bold'],
                 bg=COLORS['white'], fg=COLORS['text_dark']).pack(anchor='w', pady=(6, 2))
        notes_var = tk.StringVar()
        tk.Entry(form, textvariable=notes_var, font=FONTS['input'],
                 relief='solid', bd=1, bg=COLORS['bg_light']).pack(fill='x', ipady=5)

        lunas_btn = tk.Button(
            form, text=f"  Bayar Lunas ({format_currency(customer['debt_amount'])})  ",
            font=FONTS['small_bold'], bg=COLORS['success_bg'], fg=COLORS['success'],
            relief='flat', cursor='hand2',
            command=lambda: amount_var.set(str(int(customer['debt_amount']))),
            pady=5
        )
        lunas_btn.pack(fill='x', pady=(8, 0))

        def pay():
            try:
                amount = float(amount_var.get().replace('.', '').replace(',', ''))
                if amount <= 0:
                    messagebox.showwarning("Peringatan", "Jumlah bayar harus lebih dari 0!", parent=dialog)
                    return
                if amount > customer['debt_amount']:
                    messagebox.showwarning("Peringatan",
                                           f"Jumlah melebihi hutang! Max: {format_currency(customer['debt_amount'])}",
                                           parent=dialog)
                    return

                db.add_debt_payment(customer_id, amount, self.user['id'], notes_var.get())
                messagebox.showinfo("Berhasil",
                                     f"Pembayaran hutang {format_currency(amount)} berhasil dicatat!",
                                     parent=dialog)
                dialog.destroy()
                self.refresh_data()
            except ValueError:
                messagebox.showerror("Error", "Masukkan jumlah yang valid!", parent=dialog)

        pay_btn = tk.Button(
            btn_frame, text="  💾  Simpan Pembayaran  ", font=FONTS['button'],
            bg=COLORS['success'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=pay, padx=20, pady=8
        )
        pay_btn.pack(side='left', padx=(0, 10))
        pay_btn.bind('<Enter>', lambda e: pay_btn.config(bg='#1B5E20'))
        pay_btn.bind('<Leave>', lambda e: pay_btn.config(bg=COLORS['success']))

        cancel_btn = tk.Button(
            btn_frame, text="  Batal  ", font=FONTS['button'],
            bg=COLORS['border'], fg=COLORS['text_dark'], relief='flat',
            cursor='hand2', command=dialog.destroy, padx=20, pady=8
        )
        cancel_btn.pack(side='left')

    def _show_debt_history(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pelanggan!")
            return

        values = self.tree.item(selected[0])['values']
        customer_id = values[0]
        customer_name = values[1]
        payments = db.get_debt_payments(customer_id)

        dialog = tk.Toplevel(self)
        dialog.title(f"Riwayat Hutang - {customer_name}")
        dialog.geometry("640x500")
        dialog.minsize(550, 420)
        dialog.resizable(True, True)
        dialog.configure(bg=COLORS['white'])
        dialog.transient(self)
        dialog.grab_set()

        dialog.update_idletasks()
        x = (dialog.winfo_screenwidth() // 2) - 320
        y = (dialog.winfo_screenheight() // 2) - 250
        dialog.geometry(f'+{x}+{y}')

        tk.Label(dialog, text="Riwayat Pembayaran Hutang", font=FONTS['subtitle'],
                 bg=COLORS['white'], fg=COLORS['primary_dark']).pack(side='top', pady=(15, 5))
        tk.Label(dialog, text=f"Pelanggan: {customer_name}", font=FONTS['body'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(side='top', pady=(0, 10))

        customer = db.get_customer_by_id(customer_id)
        debt_label = tk.Label(
            dialog, text=f"Sisa Hutang: {format_currency(customer['debt_amount'])}",
            font=FONTS['heading'], bg=COLORS['white'],
            fg=COLORS['danger'] if customer['debt_amount'] > 0 else COLORS['success']
        )
        debt_label.pack(side='top', pady=(0, 10))

        sep = tk.Frame(dialog, bg=COLORS['border_light'], height=1)
        sep.pack(side='top', fill='x', padx=20)

        # Buttons packed at BOTTOM first so they are ALWAYS visible
        btn_box = tk.Frame(dialog, bg=COLORS['white'])
        btn_box.pack(side='bottom', fill='x', padx=20, pady=15)

        if payments:
            table_frame = tk.Frame(dialog, bg=COLORS['white'])
            table_frame.pack(side='top', fill='both', expand=True, padx=20, pady=10)

            cols = ('date', 'amount', 'user', 'notes')
            tree = ttk.Treeview(table_frame, columns=cols, show='headings', height=10)
            tree.heading('date', text='Tanggal')
            tree.heading('amount', text='Jumlah')
            tree.heading('user', text='Dicatat oleh')
            tree.heading('notes', text='Catatan')

            tree.column('date', width=140, anchor='center')
            tree.column('amount', width=130, anchor='e')
            tree.column('user', width=120, anchor='center')
            tree.column('notes', width=140, anchor='w')

            for p in payments:
                tree.insert('', 'end', values=(
                    p['created_at'][:19] if p['created_at'] else '',
                    format_currency(p['amount']),
                    p['user_name'],
                    p['notes'] or '-'
                ))
            tree.pack(fill='both', expand=True)
        else:
            tk.Label(dialog, text="Belum ada riwayat pembayaran", font=FONTS['body'],
                     bg=COLORS['white'], fg=COLORS['text_light']).pack(side='top', pady=30)

        def export_to_excel():
            if not payments:
                messagebox.showinfo("Info", "Tidak ada data riwayat untuk diunduh.", parent=dialog)
                return
            file_path = filedialog.asksaveasfilename(
                parent=dialog,
                title="Unduh Riwayat Hutang",
                defaultextension=".xlsx",
                initialfile=f"Riwayat_Hutang_{customer_name}.xlsx",
                filetypes=[("Excel Workbook (*.xlsx)", "*.xlsx"), ("All Files (*.*)", "*.*")]
            )
            if not file_path:
                return
            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Riwayat Hutang"
                ws.append(["RIWAYAT PEMBAYARAN HUTANG"])
                ws.append([f"Pelanggan: {customer_name}"])
                ws.append([f"Sisa Hutang Saat Ini: {format_currency(customer['debt_amount'])}"])
                ws.append([])
                headers = ["No", "Tanggal & Waktu", "Jumlah Bayar (Rp)", "Dicatat Oleh", "Catatan"]
                ws.append(headers)
                header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
                header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=5, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center" if col_idx != 3 and col_idx != 5 else "right")

                total_paid = 0
                for i, p in enumerate(payments, 1):
                    total_paid += p['amount']
                    ws.append([
                        i,
                        p['created_at'][:19] if p['created_at'] else '',
                        p['amount'],
                        p['user_name'],
                        p['notes'] or '-'
                    ])
                ws.append([])
                ws.append(["", "TOTAL DIBAYAR", total_paid, "", ""])
                ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
                ws.cell(row=ws.max_row, column=3).font = Font(bold=True)

                for col in ws.columns:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = openpyxl.utils.get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

                wb.save(file_path)
                messagebox.showinfo("Berhasil", f"Data berhasil diunduh ke:\n{file_path}", parent=dialog)
            except Exception as e:
                messagebox.showerror("Error", f"Gagal mengunduh Excel:\n{e}", parent=dialog)

        export_btn = tk.Button(
            btn_box, text="  💾  Unduh ke Excel (.xlsx)  ", font=FONTS['button'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=export_to_excel, padx=20, pady=8
        )
        export_btn.pack(side='left', padx=(0, 10))

        close_btn = tk.Button(
            btn_box, text="  Tutup  ", font=FONTS['button'],
            bg=COLORS['border'], fg=COLORS['text_dark'], relief='flat',
            cursor='hand2', command=dialog.destroy, padx=20, pady=8
        )
        close_btn.pack(side='left')

    def _delete_customer(self):
        selected = self.tree.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih pelanggan yang ingin dihapus!")
            return

        values = self.tree.item(selected[0])['values']
        if messagebox.askyesno("Konfirmasi", f"Hapus pelanggan '{values[1]}'?"):
            db.delete_customer(values[0])
            messagebox.showinfo("Berhasil", "Pelanggan berhasil dihapus!")
            self.refresh_data()
