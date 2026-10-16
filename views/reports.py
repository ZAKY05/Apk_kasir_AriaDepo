"""
Reports View - Laporan Keuangan (Pemasukan & Pengeluaran)
Dilengkapi dengan tabel pemasukan, tabel pengeluaran, popup detail struk dengan tombol simpan,
serta fitur unduh laporan ke format Excel (.xlsx).
"""

import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime
import csv

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from styles import COLORS, FONTS, ICONS, format_currency, format_number
import database as db

CAT_LABELS = {
    'supplier': '🚰 Beli Air Supplier',
    'peralatan': '🔧 Peralatan',
    'lainnya': '📎 Lainnya',
}


class ReportsView(tk.Frame):
    def __init__(self, parent, user):
        super().__init__(parent, bg=COLORS['bg_light'])
        self.user = user
        self.current_report = None
        self._build_ui()
        self._load_report()

    def _build_ui(self):
        # Scrollable container
        canvas = tk.Canvas(self, bg=COLORS['bg_light'], highlightthickness=0)
        scrollbar = tk.Scrollbar(self, orient='vertical', command=canvas.yview)
        self.scroll_frame = tk.Frame(canvas, bg=COLORS['bg_light'])

        self.scroll_frame.bind('<Configure>', lambda e: canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.create_window((0, 0), window=self.scroll_frame, anchor='nw')
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side='left', fill='both', expand=True)
        scrollbar.pack(side='right', fill='y')
        canvas.bind_all('<MouseWheel>', lambda e: canvas.yview_scroll(int(-1 * (e.delta / 120)), 'units'))

        # Header
        header = tk.Frame(self.scroll_frame, bg=COLORS['bg_light'])
        header.pack(fill='x', padx=25, pady=(25, 15))

        tk.Label(
            header, text=f"{ICONS['report']}  Laporan Keuangan",
            font=FONTS['title'], bg=COLORS['bg_light'], fg=COLORS['text_dark']
        ).pack(side='left')

        # Tombol download laporan lengkap di header
        export_all_btn = tk.Button(
            header, text="  📥  Unduh Semua Laporan (.xlsx)  ", font=FONTS['button'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._export_full_excel, pady=8, padx=15
        )
        export_all_btn.pack(side='right')
        export_all_btn.bind('<Enter>', lambda e: export_all_btn.config(bg=COLORS['primary_hover']))
        export_all_btn.bind('<Leave>', lambda e: export_all_btn.config(bg=COLORS['primary']))

        # Period selector
        period_frame = tk.Frame(self.scroll_frame, bg=COLORS['white'], padx=15, pady=12,
                                 highlightbackground=COLORS['border_light'], highlightthickness=1)
        period_frame.pack(fill='x', padx=25, pady=(0, 15))

        tk.Label(period_frame, text="Periode:", font=FONTS['body_bold'],
                 bg=COLORS['white'], fg=COLORS['text_dark']).pack(side='left', padx=(0, 10))

        self.period_var = tk.StringVar(value="daily")

        periods = [
            ("daily", f"{ICONS['calendar']}  Harian"),
            ("weekly", "📅  Mingguan"),
            ("monthly", "📅  Bulanan"),
            ("yearly", "📅  Tahunan"),
        ]

        for val, text in periods:
            btn = tk.Radiobutton(
                period_frame, text=text, variable=self.period_var, value=val,
                font=FONTS['body'], bg=COLORS['white'], activebackground=COLORS['white'],
                selectcolor=COLORS['white'], command=self._load_report, cursor='hand2',
                indicatoron=0, padx=15, pady=6, relief='flat',
                selectimage=None
            )
            btn.pack(side='left', padx=3)

        # Custom date range
        tk.Label(period_frame, text="    Dari:", font=FONTS['body'],
                 bg=COLORS['white'], fg=COLORS['text_medium']).pack(side='left', padx=(15, 5))

        self.date_from_var = tk.StringVar()
        tk.Entry(period_frame, textvariable=self.date_from_var, font=FONTS['small'],
                 width=12, relief='solid', bd=1).pack(side='left')

        tk.Label(period_frame, text="  s/d:", font=FONTS['body'],
                 bg=COLORS['white'], fg=COLORS['text_medium']).pack(side='left', padx=5)

        self.date_to_var = tk.StringVar()
        tk.Entry(period_frame, textvariable=self.date_to_var, font=FONTS['small'],
                 width=12, relief='solid', bd=1).pack(side='left')

        search_btn = tk.Button(
            period_frame, text=f"{ICONS['search']}  Tampilkan", font=FONTS['body_bold'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._load_custom_report, padx=12, pady=4
        )
        search_btn.pack(side='left', padx=10)

        # Summary cards
        self.summary_frame = tk.Frame(self.scroll_frame, bg=COLORS['bg_light'])
        self.summary_frame.pack(fill='x', padx=25, pady=(0, 15))
        self.summary_frame.columnconfigure((0, 1, 2, 3), weight=1)

        # ==================== TABEL PEMASUKAN ====================
        self.income_frame = tk.LabelFrame(
            self.scroll_frame, text=f"  💰  Tabel Pemasukan (Penjualan)  ",
            font=FONTS['heading'], bg=COLORS['white'], fg=COLORS['text_dark'],
            relief='solid', bd=1, padx=15, pady=12
        )
        self.income_frame.pack(fill='x', padx=25, pady=(0, 18))

        # Toolbar tabel pemasukan
        income_toolbar = tk.Frame(self.income_frame, bg=COLORS['white'])
        income_toolbar.pack(fill='x', pady=(0, 10))

        self.income_count_label = tk.Label(
            income_toolbar, text="Memuat data...", font=FONTS['body_bold'],
            bg=COLORS['white'], fg=COLORS['text_dark']
        )
        self.income_count_label.pack(side='left')

        export_income_btn = tk.Button(
            income_toolbar, text="  📥  Unduh Pemasukan (.xlsx)  ", font=FONTS['small_bold'],
            bg=COLORS['success'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._export_sales_excel, padx=10, pady=5
        )
        export_income_btn.pack(side='right', padx=(6, 0))

        detail_trx_btn = tk.Button(
            income_toolbar, text="  👁️  Lihat Detail Transaksi  ", font=FONTS['small_bold'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._show_selected_transaction_detail, padx=10, pady=5
        )
        detail_trx_btn.pack(side='right')

        # Container Treeview Pemasukan
        income_table_box = tk.Frame(self.income_frame, bg=COLORS['white'])
        income_table_box.pack(fill='both', expand=True)

        income_cols = ('invoice', 'time', 'customer', 'kasir', 'method', 'discount', 'total')
        self.income_tree = ttk.Treeview(income_table_box, columns=income_cols, show='headings', height=8)

        self.income_tree.heading('invoice', text='No. Invoice')
        self.income_tree.heading('time', text='Waktu Transaksi')
        self.income_tree.heading('customer', text='Pelanggan')
        self.income_tree.heading('kasir', text='Kasir')
        self.income_tree.heading('method', text='Metode')
        self.income_tree.heading('discount', text='Diskon')
        self.income_tree.heading('total', text='Total Pemasukan')

        self.income_tree.column('invoice', width=130, anchor='center')
        self.income_tree.column('time', width=140, anchor='center')
        self.income_tree.column('customer', width=130, anchor='w')
        self.income_tree.column('kasir', width=120, anchor='center')
        self.income_tree.column('method', width=90, anchor='center')
        self.income_tree.column('discount', width=100, anchor='e')
        self.income_tree.column('total', width=130, anchor='e')

        income_scroll = ttk.Scrollbar(income_table_box, orient='vertical', command=self.income_tree.yview)
        self.income_tree.configure(yscrollcommand=income_scroll.set)

        self.income_tree.pack(side='left', fill='both', expand=True)
        income_scroll.pack(side='right', fill='y')

        # Double-click to view detail
        self.income_tree.bind('<Double-1>', lambda e: self._show_selected_transaction_detail())

        # ==================== TABEL PENGELUARAN ====================
        self.expense_table_frame = tk.LabelFrame(
            self.scroll_frame, text=f"  💸  Tabel Pengeluaran  ",
            font=FONTS['heading'], bg=COLORS['white'], fg=COLORS['text_dark'],
            relief='solid', bd=1, padx=15, pady=12
        )
        self.expense_table_frame.pack(fill='x', padx=25, pady=(0, 18))

        # Toolbar tabel pengeluaran
        expense_toolbar = tk.Frame(self.expense_table_frame, bg=COLORS['white'])
        expense_toolbar.pack(fill='x', pady=(0, 10))

        self.expense_count_label = tk.Label(
            expense_toolbar, text="Memuat data...", font=FONTS['body_bold'],
            bg=COLORS['white'], fg=COLORS['text_dark']
        )
        self.expense_count_label.pack(side='left')

        export_expense_btn = tk.Button(
            expense_toolbar, text="  📥  Unduh Pengeluaran (.xlsx)  ", font=FONTS['small_bold'],
            bg=COLORS['danger'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=self._export_expenses_excel, padx=10, pady=5
        )
        export_expense_btn.pack(side='right')

        # Container Treeview Pengeluaran
        expense_table_box = tk.Frame(self.expense_table_frame, bg=COLORS['white'])
        expense_table_box.pack(fill='both', expand=True)

        exp_cols = ('no', 'date', 'category', 'description', 'amount', 'notes', 'user')
        self.expense_tree = ttk.Treeview(expense_table_box, columns=exp_cols, show='headings', height=8)

        self.expense_tree.heading('no', text='No')
        self.expense_tree.heading('date', text='Tanggal')
        self.expense_tree.heading('category', text='Kategori')
        self.expense_tree.heading('description', text='Deskripsi Pengeluaran')
        self.expense_tree.heading('amount', text='Jumlah (Rp)')
        self.expense_tree.heading('notes', text='Catatan')
        self.expense_tree.heading('user', text='Dicatat Oleh')

        self.expense_tree.column('no', width=45, anchor='center')
        self.expense_tree.column('date', width=110, anchor='center')
        self.expense_tree.column('category', width=140, anchor='w')
        self.expense_tree.column('description', width=180, anchor='w')
        self.expense_tree.column('amount', width=120, anchor='e')
        self.expense_tree.column('notes', width=140, anchor='w')
        self.expense_tree.column('user', width=110, anchor='center')

        exp_scroll = ttk.Scrollbar(expense_table_box, orient='vertical', command=self.expense_tree.yview)
        self.expense_tree.configure(yscrollcommand=exp_scroll.set)

        self.expense_tree.pack(side='left', fill='both', expand=True)
        exp_scroll.pack(side='right', fill='y')

        # Top products & Expense breakdown
        bottom = tk.Frame(self.scroll_frame, bg=COLORS['bg_light'])
        bottom.pack(fill='x', padx=25, pady=(0, 25))
        bottom.columnconfigure(0, weight=1)
        bottom.columnconfigure(1, weight=1)

        self.top_products_frame = tk.LabelFrame(
            bottom, text=f"  {ICONS['product']}  Produk Terlaris  ",
            font=FONTS['heading'], bg=COLORS['white'], fg=COLORS['text_dark'],
            relief='solid', bd=1, padx=15, pady=15
        )
        self.top_products_frame.grid(row=0, column=0, sticky='nsew', padx=(0, 8))

        self.expense_category_frame = tk.LabelFrame(
            bottom, text=f"  {ICONS['expense']}  Pengeluaran per Kategori  ",
            font=FONTS['heading'], bg=COLORS['white'], fg=COLORS['text_dark'],
            relief='solid', bd=1, padx=15, pady=15
        )
        self.expense_category_frame.grid(row=0, column=1, sticky='nsew')

    def _create_summary_card(self, parent, title, value, color, col):
        """Buat kartu summary"""
        card = tk.Frame(parent, bg=COLORS['white'], padx=18, pady=14,
                        highlightbackground=COLORS['border_light'], highlightthickness=1)
        card.grid(row=0, column=col, sticky='nsew', padx=4, pady=4)

        tk.Label(card, text=title, font=FONTS['small_bold'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(anchor='w')
        tk.Label(card, text=value, font=FONTS['medium_number'],
                 bg=COLORS['white'], fg=color).pack(anchor='w', pady=(5, 0))

    def _load_report(self):
        """Load laporan berdasarkan periode"""
        period = self.period_var.get()
        report = db.get_sales_report(period)

        self.date_from_var.set(report['start_date'])
        self.date_to_var.set(report['end_date'])

        self._display_report(report)

    def _load_custom_report(self):
        """Load laporan dengan tanggal kustom"""
        start = self.date_from_var.get().strip()
        end = self.date_to_var.get().strip()

        if not start or not end:
            messagebox.showwarning("Peringatan", "Tanggal awal dan akhir harus diisi!")
            return

        report = db.get_sales_report('daily', start, end)
        self._display_report(report)

    def _display_report(self, report):
        """Tampilkan data laporan ke UI"""
        self.current_report = report

        # Clear summary cards
        for widget in self.summary_frame.winfo_children():
            widget.destroy()

        # Summary cards
        self._create_summary_card(
            self.summary_frame, "Total Pemasukan",
            format_currency(report['total_sales']), COLORS['primary'], 0
        )
        self._create_summary_card(
            self.summary_frame, "Jumlah Transaksi",
            format_number(report['total_count']), COLORS['success'], 1
        )
        self._create_summary_card(
            self.summary_frame, "Total Pengeluaran",
            format_currency(report['total_expenses']), COLORS['danger'], 2
        )
        net_color = COLORS['success'] if report['net_profit'] >= 0 else COLORS['danger']
        self._create_summary_card(
            self.summary_frame, "Laba Bersih",
            format_currency(report['net_profit']), net_color, 3
        )

        # Populate Tabel Pemasukan
        for item in self.income_tree.get_children():
            self.income_tree.delete(item)

        transactions = report.get('transactions', [])
        total_pemasukan = sum(t['total'] for t in transactions)
        self.income_count_label.config(
            text=f"Total: {len(transactions)} transaksi | {format_currency(total_pemasukan)}"
        )

        for trx in transactions:
            self.income_tree.insert('', 'end', values=(
                trx['invoice_number'],
                trx['created_at'][:19] if trx['created_at'] else '',
                trx['customer_name'],
                trx['kasir_name'],
                trx['payment_method'].upper(),
                format_currency(trx.get('discount', 0)),
                format_currency(trx['total']),
            ))

        # Populate Tabel Pengeluaran
        for item in self.expense_tree.get_children():
            self.expense_tree.delete(item)

        expenses = report.get('expenses', [])
        total_pengeluaran = sum(e['amount'] for e in expenses)
        self.expense_count_label.config(
            text=f"Total: {len(expenses)} pengeluaran | {format_currency(total_pengeluaran)}"
        )

        for i, exp in enumerate(expenses, 1):
            cat_display = CAT_LABELS.get(exp['category'], exp['category'])
            self.expense_tree.insert('', 'end', values=(
                i,
                exp['expense_date'],
                cat_display,
                exp['description'],
                format_currency(exp['amount']),
                exp['notes'] or '-',
                exp.get('user_name', '-')
            ))

        # Top products
        for widget in self.top_products_frame.winfo_children():
            widget.destroy()

        if report.get('top_products'):
            for i, prod in enumerate(report['top_products'], 1):
                row = tk.Frame(self.top_products_frame, bg=COLORS['white'])
                row.pack(fill='x', pady=3)

                badge_color = COLORS['primary'] if i <= 3 else COLORS['text_light']
                tk.Label(row, text=f"  #{i}", font=FONTS['body_bold'],
                         bg=COLORS['white'], fg=badge_color, width=5).pack(side='left')
                tk.Label(row, text=prod['product_name'], font=FONTS['body'],
                         bg=COLORS['white'], fg=COLORS['text_dark']).pack(side='left')
                tk.Label(row, text=f"{prod['total_qty']} terjual", font=FONTS['small'],
                         bg=COLORS['white'], fg=COLORS['text_light']).pack(side='right', padx=(0, 10))
                tk.Label(row, text=format_currency(prod['total_revenue']), font=FONTS['body_bold'],
                         bg=COLORS['white'], fg=COLORS['success']).pack(side='right')
        else:
            tk.Label(self.top_products_frame, text="Belum ada data",
                     font=FONTS['body'], bg=COLORS['white'],
                     fg=COLORS['text_light']).pack(pady=20)

        # Expense breakdown per kategori
        for widget in self.expense_category_frame.winfo_children():
            widget.destroy()

        if report.get('expense_by_category'):
            for exp in report['expense_by_category']:
                row = tk.Frame(self.expense_category_frame, bg=COLORS['white'])
                row.pack(fill='x', pady=4)
                tk.Label(row, text=CAT_LABELS.get(exp['category'], exp['category']),
                         font=FONTS['body'], bg=COLORS['white'],
                         fg=COLORS['text_dark']).pack(side='left')
                tk.Label(row, text=format_currency(exp['total']), font=FONTS['body_bold'],
                         bg=COLORS['white'], fg=COLORS['danger']).pack(side='right')
        else:
            tk.Label(self.expense_category_frame, text="Tidak ada pengeluaran",
                     font=FONTS['body'], bg=COLORS['white'],
                     fg=COLORS['text_light']).pack(pady=20)

    def _show_selected_transaction_detail(self):
        """Buka pop up detail transaksi dengan tombol Simpan Struk"""
        selected = self.income_tree.selection()
        if not selected:
            messagebox.showwarning("Peringatan", "Pilih transaksi dari tabel pemasukan terlebih dahulu!")
            return

        values = self.income_tree.item(selected[0])['values']
        invoice_number = str(values[0])

        trx = db.get_transaction_by_invoice(invoice_number)
        if not trx:
            messagebox.showerror("Error", f"Data transaksi {invoice_number} tidak ditemukan!")
            return

        dialog = tk.Toplevel(self)
        dialog.title(f"Detail Transaksi - {invoice_number}")
        dialog.geometry("480x600")
        dialog.minsize(420, 500)
        dialog.resizable(True, True)
        dialog.configure(bg=COLORS['white'])
        dialog.transient(self)
        dialog.grab_set()

        dialog.update_idletasks()
        x = (dialog.winfo_screenwidth() // 2) - 240
        y = (dialog.winfo_screenheight() // 2) - 300
        dialog.geometry(f'+{x}+{y}')

        # Header dialog
        tk.Label(dialog, text="🧾", font=('Segoe UI', 32), bg=COLORS['white']).pack(side='top', pady=(15, 2))
        tk.Label(dialog, text="Detail Transaksi Penjualan", font=FONTS['subtitle'],
                 bg=COLORS['white'], fg=COLORS['primary_dark']).pack(side='top')

        sep = tk.Frame(dialog, bg=COLORS['border_light'], height=1)
        sep.pack(side='top', fill='x', padx=25, pady=10)

        # Buttons packed at BOTTOM first so they are ALWAYS visible
        btn_box = tk.Frame(dialog, bg=COLORS['white'])
        btn_box.pack(side='bottom', fill='x', padx=25, pady=(10, 15))

        # Info Ringkas
        info_frame = tk.Frame(dialog, bg=COLORS['white'], padx=25)
        info_frame.pack(side='top', fill='x')

        def add_info_row(label, value, bold=False):
            r = tk.Frame(info_frame, bg=COLORS['white'])
            r.pack(fill='x', pady=1)
            tk.Label(r, text=label, font=FONTS['body'], bg=COLORS['white'], fg=COLORS['text_medium']).pack(side='left')
            f = FONTS['body_bold'] if bold else FONTS['body']
            tk.Label(r, text=value, font=f, bg=COLORS['white'], fg=COLORS['text_dark']).pack(side='right')

        add_info_row("No. Invoice:", trx['invoice_number'], True)
        add_info_row("Waktu:", trx['created_at'][:19] if trx['created_at'] else '-')
        add_info_row("Pelanggan:", trx['customer_name'])
        add_info_row("Kasir:", trx['kasir_name'])
        add_info_row("Metode Pembayaran:", trx['payment_method'].upper())
        if trx.get('is_delivery'):
            add_info_row("Kurir:", trx.get('courier_name') or '-')
            add_info_row("Ongkir Kurir:", format_currency(trx.get('delivery_fee', 0)))
        if trx.get('discount', 0) > 0:
            add_info_row("Diskon:", format_currency(trx['discount']))
        add_info_row("Total Transaksi:", format_currency(trx['total']), True)

        sep2 = tk.Frame(dialog, bg=COLORS['border_light'], height=1)
        sep2.pack(side='top', fill='x', padx=25, pady=8)

        # Item-item transaksi
        tk.Label(dialog, text="Item yang Dibeli:", font=FONTS['small_bold'],
                 bg=COLORS['white'], fg=COLORS['text_light']).pack(side='top', anchor='w', padx=25)

        items_box = tk.Frame(dialog, bg=COLORS['bg_light'], padx=10, pady=8,
                             highlightbackground=COLORS['border_light'], highlightthickness=1)
        items_box.pack(side='top', fill='both', expand=True, padx=25, pady=(5, 5))

        canvas = tk.Canvas(items_box, bg=COLORS['bg_light'], highlightthickness=0)
        scrollbar = tk.Scrollbar(items_box, orient='vertical', command=canvas.yview)
        scrollable_frame = tk.Frame(canvas, bg=COLORS['bg_light'])

        scrollable_frame.bind('<Configure>', lambda e: canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.create_window((0, 0), window=scrollable_frame, anchor='nw')
        canvas.configure(yscrollcommand=scrollbar.set)

        canvas.pack(side='left', fill='both', expand=True)
        scrollbar.pack(side='right', fill='y')

        items = trx.get('items', [])
        for item in items:
            row = tk.Frame(scrollable_frame, bg=COLORS['bg_light'])
            row.pack(fill='x', pady=2)
            tk.Label(row, text=f"• {item['product_name']} x{item['quantity']}",
                     font=FONTS['body'], bg=COLORS['bg_light'], fg=COLORS['text_dark']).pack(side='left')
            tk.Label(row, text=format_currency(item['subtotal']),
                     font=FONTS['body_bold'], bg=COLORS['bg_light'], fg=COLORS['primary']).pack(side='right')

        def save_receipt_file():
            file_path = filedialog.asksaveasfilename(
                parent=dialog,
                title="Simpan Struk Transaksi",
                defaultextension=".txt",
                initialfile=f"Struk_{trx['invoice_number']}.txt",
                filetypes=[("Text File (*.txt)", "*.txt"), ("All Files (*.*)", "*.*")]
            )
            if not file_path:
                return
            try:
                lines = [
                    "=" * 42,
                    "              DEPO AIR MINUM              ",
                    "=" * 42,
                    f"No. Invoice : {trx['invoice_number']}",
                    f"Tanggal     : {trx['created_at'][:19] if trx['created_at'] else '-'}",
                    f"Pelanggan   : {trx['customer_name']}",
                    f"Kasir       : {trx['kasir_name']}",
                    f"Metode      : {trx['payment_method'].upper()}",
                ]
                if trx.get('is_delivery'):
                    lines.append(f"Kurir       : {trx.get('courier_name') or '-'}")
                    lines.append(f"Ongkir      : {format_currency(trx.get('delivery_fee', 0))}")
                lines.append("-" * 42)
                lines.append(f"{'Produk':<24}{'Qty':>4}{'Subtotal':>14}")
                lines.append("-" * 42)
                for itm in items:
                    p_name = itm.get('product_name', 'Produk')
                    p_qty = itm.get('quantity', 1)
                    p_sub = itm.get('subtotal', 0)
                    lines.append(f"{p_name[:22]:<24}{p_qty:>4}{format_currency(p_sub):>14}")
                lines.append("-" * 42)
                if trx.get('discount', 0) > 0:
                    lines.append(f"{'Diskon':<24}{format_currency(trx['discount']):>18}")
                lines.append(f"{'TOTAL':<24}{format_currency(trx['total']):>18}")
                if trx.get('paid_amount', 0) > 0:
                    lines.append(f"{'DIBAYAR':<24}{format_currency(trx['paid_amount']):>18}")
                    lines.append(f"{'KEMBALIAN':<24}{format_currency(trx.get('change_amount', 0)):>18}")
                lines.append("=" * 42)
                lines.append("       Terima Kasih Atas Kunjungan Anda!   ")
                lines.append("=" * 42)

                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write("\n".join(lines) + "\n")

                messagebox.showinfo("Berhasil", f"Struk berhasil disimpan di:\n{file_path}", parent=dialog)
            except Exception as e:
                messagebox.showerror("Error", f"Gagal menyimpan struk:\n{e}", parent=dialog)

        save_btn = tk.Button(
            btn_box, text="  💾  Simpan Struk  ", font=FONTS['button'],
            bg=COLORS['primary'], fg=COLORS['text_white'], relief='flat',
            cursor='hand2', command=save_receipt_file, pady=9
        )
        save_btn.pack(side='left', fill='x', expand=True, padx=(0, 10))

        close_btn = tk.Button(
            btn_box, text="Tutup", font=FONTS['button'],
            bg=COLORS['border'], fg=COLORS['text_dark'], relief='flat',
            cursor='hand2', command=dialog.destroy, pady=9
        )
        close_btn.pack(side='right', fill='x', expand=True)

    # ==================== EXCEL EXPORT FUNCTIONS ====================

    def _apply_sheet_styling(self, ws, header_row=1):
        """Beri gaya rapi ke sheet openpyxl (header, font, border, autofit)"""
        header_fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        thin_side = Side(border_style="thin", color="D5D8DC")
        cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = cell_border

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if '\n' in val:
                    val = max(val.split('\n'), key=len)
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    def _export_full_excel(self):
        """Unduh Laporan Lengkap (Ringkasan, Pemasukan, Pengeluaran) ke Excel"""
        if not self.current_report:
            messagebox.showwarning("Peringatan", "Tidak ada data laporan untuk diunduh!", parent=self)
            return

        if not HAS_OPENPYXL:
            messagebox.showwarning(
                "Peringatan",
                "Library 'openpyxl' belum terpasang.\nSilakan jalankan: pip install openpyxl\nuntuk mengaktifkan fitur ekspor Excel.",
                parent=self
            )
            return

        report = self.current_report
        start = report['start_date']
        end = report['end_date']

        file_path = filedialog.asksaveasfilename(
            parent=self,
            title="Unduh Laporan Keuangan Lengkap",
            defaultextension=".xlsx",
            initialfile=f"Laporan_Keuangan_{start}_sd_{end}.xlsx",
            filetypes=[("Excel Workbook (*.xlsx)", "*.xlsx"), ("All Files (*.*)", "*.*")]
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()

            # ===== Sheet 1: Ringkasan =====
            ws_summary = wb.active
            ws_summary.title = "Ringkasan"

            ws_summary.append(["LAPORAN KEUANGAN DEPO AIR"])
            ws_summary.append([f"Periode: {start} s/d {end}"])
            ws_summary.append([f"Waktu Ekspor: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"])
            ws_summary.append([])

            ws_summary.append(["RINGKASAN KEUANGAN", "NILAI"])
            ws_summary.append(["Total Pemasukan (Penjualan)", report['total_sales']])
            ws_summary.append(["Jumlah Transaksi", report['total_count']])
            ws_summary.append(["Total Pengeluaran", report['total_expenses']])
            ws_summary.append(["Laba Bersih", report['net_profit']])
            ws_summary.append([])

            # Header ringkasan
            for col_idx in (1, 2):
                c = ws_summary.cell(row=5, column=col_idx)
                c.fill = PatternFill(start_color="1A5276", end_color="1A5276", fill_type="solid")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

            # Produk terlaris
            ws_summary.append(["PRODUK TERLARIS (TOP 10)", "QTY TERJUAL", "TOTAL REVENUE (RP)"])
            top_header_row = ws_summary.max_row
            for col_idx in (1, 2, 3):
                c = ws_summary.cell(row=top_header_row, column=col_idx)
                c.fill = PatternFill(start_color="2874A6", end_color="2874A6", fill_type="solid")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

            for prod in report.get('top_products', []):
                ws_summary.append([prod['product_name'], prod['total_qty'], prod['total_revenue']])

            ws_summary.append([])

            # Kategori Pengeluaran
            ws_summary.append(["PENGELUARAN PER KATEGORI", "TOTAL (RP)"])
            cat_header_row = ws_summary.max_row
            for col_idx in (1, 2):
                c = ws_summary.cell(row=cat_header_row, column=col_idx)
                c.fill = PatternFill(start_color="922B21", end_color="922B21", fill_type="solid")
                c.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

            for cat in report.get('expense_by_category', []):
                cat_label = CAT_LABELS.get(cat['category'], cat['category'])
                ws_summary.append([cat_label, cat['total']])

            for col in ws_summary.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 15)

            # ===== Sheet 2: Pemasukan =====
            ws_income = wb.create_sheet(title="Pemasukan")
            income_headers = ["No. Invoice", "Tanggal & Waktu", "Pelanggan", "Kasir", "Metode Bayar", "Diskon (Rp)", "Total (Rp)"]
            ws_income.append(income_headers)

            total_inc = 0
            for trx in report.get('transactions', []):
                total_inc += trx['total']
                ws_income.append([
                    trx['invoice_number'],
                    trx['created_at'][:19] if trx['created_at'] else '',
                    trx['customer_name'],
                    trx['kasir_name'],
                    trx['payment_method'].upper(),
                    trx.get('discount', 0),
                    trx['total']
                ])

            ws_income.append([])
            ws_income.append(["", "", "", "", "TOTAL PEMASUKAN", "", total_inc])
            ws_income.cell(row=ws_income.max_row, column=5).font = Font(name="Segoe UI", bold=True)
            ws_income.cell(row=ws_income.max_row, column=7).font = Font(name="Segoe UI", bold=True)
            self._apply_sheet_styling(ws_income, header_row=1)

            # ===== Sheet 3: Pengeluaran =====
            ws_expense = wb.create_sheet(title="Pengeluaran")
            expense_headers = ["No", "Tanggal", "Kategori", "Deskripsi Pengeluaran", "Jumlah (Rp)", "Catatan", "Dicatat Oleh"]
            ws_expense.append(expense_headers)

            total_exp = 0
            for i, exp in enumerate(report.get('expenses', []), 1):
                total_exp += exp['amount']
                ws_expense.append([
                    i,
                    exp['expense_date'],
                    CAT_LABELS.get(exp['category'], exp['category']),
                    exp['description'],
                    exp['amount'],
                    exp['notes'] or '-',
                    exp.get('user_name', '-')
                ])

            ws_expense.append([])
            ws_expense.append(["", "", "", "TOTAL PENGELUARAN", total_exp, "", ""])
            ws_expense.cell(row=ws_expense.max_row, column=4).font = Font(name="Segoe UI", bold=True)
            ws_expense.cell(row=ws_expense.max_row, column=5).font = Font(name="Segoe UI", bold=True)
            self._apply_sheet_styling(ws_expense, header_row=1)

            # ===== Sheet 4: Flow Meter Air =====
            ws_meters = wb.create_sheet(title="Flow Meter Air")
            meter_headers = [
                "No", "Tanggal", "Shift", "Petugas",
                "Bio Awal", "Bio Akhir", "Bio (L)", "Bio (Galon)", "Bio (Rp)",
                "RO Awal", "RO Akhir", "RO (L)", "RO (Galon)", "RO (Rp)",
                "Total Liter", "Total Galon", "Total Omzet (Rp)", "Catatan"
            ]
            ws_meters.append(meter_headers)
            water_records = db.get_water_meters(start, end, limit=500)
            tot_w_lit = 0
            tot_w_gal = 0
            tot_w_rev = 0
            for i, wm in enumerate(reversed(water_records), 1):
                tot_w_lit += wm['total_liters']
                tot_w_gal += wm['total_gallons']
                tot_w_rev += wm['total_revenue']
                ws_meters.append([
                    i, wm['record_date'], wm['shift_name'], wm['user_name'] or '-',
                    wm['bio_meter_start'], wm['bio_meter_end'], wm['bio_liters'], wm['bio_gallons'], wm['bio_revenue'],
                    wm['ro_meter_start'], wm['ro_meter_end'], wm['ro_liters'], wm['ro_gallons'], wm['ro_revenue'],
                    wm['total_liters'], wm['total_gallons'], wm['total_revenue'], wm['notes'] or '-'
                ])
            ws_meters.append([])
            ws_meters.append(["", "TOTAL", "", "", "", "", "", "", "", "", "", "", "", "", tot_w_lit, tot_w_gal, tot_w_rev, ""])
            ws_meters.cell(row=ws_meters.max_row, column=2).font = Font(name="Segoe UI", bold=True)
            ws_meters.cell(row=ws_meters.max_row, column=15).font = Font(name="Segoe UI", bold=True)
            ws_meters.cell(row=ws_meters.max_row, column=16).font = Font(name="Segoe UI", bold=True)
            ws_meters.cell(row=ws_meters.max_row, column=17).font = Font(name="Segoe UI", bold=True)
            self._apply_sheet_styling(ws_meters, header_row=1)

            wb.save(file_path)
            messagebox.showinfo("Berhasil", f"Laporan lengkap berhasil diunduh ke:\n{file_path}", parent=self)

        except Exception as e:
            messagebox.showerror("Error", f"Gagal mengekspor laporan ke Excel:\n{e}", parent=self)

    def _export_sales_excel(self):
        """Unduh Tabel Pemasukan ke file Excel"""
        if not self.current_report:
            messagebox.showwarning("Peringatan", "Tidak ada data pemasukan untuk diunduh!", parent=self)
            return

        if not HAS_OPENPYXL:
            messagebox.showwarning(
                "Peringatan",
                "Library 'openpyxl' belum terpasang.\nSilakan jalankan: pip install openpyxl\nuntuk mengaktifkan fitur ekspor Excel.",
                parent=self
            )
            return

        report = self.current_report
        start = report['start_date']
        end = report['end_date']

        file_path = filedialog.asksaveasfilename(
            parent=self,
            title="Unduh Tabel Pemasukan",
            defaultextension=".xlsx",
            initialfile=f"Pemasukan_{start}_sd_{end}.xlsx",
            filetypes=[("Excel Workbook (*.xlsx)", "*.xlsx"), ("All Files (*.*)", "*.*")]
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Pemasukan"

            headers = ["No. Invoice", "Tanggal & Waktu", "Pelanggan", "Kasir", "Metode Bayar", "Diskon (Rp)", "Total Pemasukan (Rp)"]
            ws.append(headers)

            total_val = 0
            for trx in report.get('transactions', []):
                total_val += trx['total']
                ws.append([
                    trx['invoice_number'],
                    trx['created_at'][:19] if trx['created_at'] else '',
                    trx['customer_name'],
                    trx['kasir_name'],
                    trx['payment_method'].upper(),
                    trx.get('discount', 0),
                    trx['total']
                ])

            ws.append([])
            ws.append(["", "", "", "", "TOTAL", "", total_val])
            ws.cell(row=ws.max_row, column=5).font = Font(name="Segoe UI", bold=True)
            ws.cell(row=ws.max_row, column=7).font = Font(name="Segoe UI", bold=True)

            self._apply_sheet_styling(ws, header_row=1)
            wb.save(file_path)
            messagebox.showinfo("Berhasil", f"Tabel pemasukan berhasil diunduh ke:\n{file_path}", parent=self)

        except Exception as e:
            messagebox.showerror("Error", f"Gagal mengunduh Excel:\n{e}", parent=self)

    def _export_expenses_excel(self):
        """Unduh Tabel Pengeluaran ke file Excel"""
        if not self.current_report:
            messagebox.showwarning("Peringatan", "Tidak ada data pengeluaran untuk diunduh!", parent=self)
            return

        if not HAS_OPENPYXL:
            messagebox.showwarning(
                "Peringatan",
                "Library 'openpyxl' belum terpasang.\nSilakan jalankan: pip install openpyxl\nuntuk mengaktifkan fitur ekspor Excel.",
                parent=self
            )
            return

        report = self.current_report
        start = report['start_date']
        end = report['end_date']

        file_path = filedialog.asksaveasfilename(
            parent=self,
            title="Unduh Tabel Pengeluaran",
            defaultextension=".xlsx",
            initialfile=f"Pengeluaran_{start}_sd_{end}.xlsx",
            filetypes=[("Excel Workbook (*.xlsx)", "*.xlsx"), ("All Files (*.*)", "*.*")]
        )
        if not file_path:
            return

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Pengeluaran"

            headers = ["No", "Tanggal", "Kategori", "Deskripsi Pengeluaran", "Jumlah (Rp)", "Catatan", "Dicatat Oleh"]
            ws.append(headers)

            total_val = 0
            for i, exp in enumerate(report.get('expenses', []), 1):
                total_val += exp['amount']
                ws.append([
                    i,
                    exp['expense_date'],
                    CAT_LABELS.get(exp['category'], exp['category']),
                    exp['description'],
                    exp['amount'],
                    exp['notes'] or '-',
                    exp.get('user_name', '-')
                ])

            ws.append([])
            ws.append(["", "", "", "TOTAL", total_val, "", ""])
            ws.cell(row=ws.max_row, column=4).font = Font(name="Segoe UI", bold=True)
            ws.cell(row=ws.max_row, column=5).font = Font(name="Segoe UI", bold=True)

            self._apply_sheet_styling(ws, header_row=1)
            wb.save(file_path)
            messagebox.showinfo("Berhasil", f"Tabel pengeluaran berhasil diunduh ke:\n{file_path}", parent=self)

        except Exception as e:
            messagebox.showerror("Error", f"Gagal mengunduh Excel:\n{e}", parent=self)
